"""Verify FIRST REVIEW FINAL.pptx: footers, literature, work plan, typos."""
import re

from pptx import Presentation

PPTX = r"C:\Users\athul\Downloads\FIRST REVIEW FINAL.pptx"
prs = Presentation(PPTX)
slides = list(prs.slides)
issues = []


def texts(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    out.append(cell.text)
    return "\n".join(out)


# 1. Footer checks -------------------------------------------------------
for pos in range(2, 19):
    s_texts = texts(slides[pos - 1])
    m = re.findall(r"Slide\. No:?\s*(\d+)", s_texts)
    if not m:
        issues.append(f"slide {pos}: NO footer number found")
        continue
    shown = int(m[-1])
    if shown != pos - 1:
        issues.append(f"slide {pos}: footer shows {shown}, expected {pos - 1}")
    # format consistency for text-box footers
    for sh in slides[pos - 1].shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text
            if "Date" in t and "IESCE" in t and "Slide. No" in t:
                if t != t.rstrip() or "No:9" in t or "  IESCE" not in t.replace("Date ", "Date  ", 0):
                    pass  # detailed check below
                if t.endswith(" "):
                    issues.append(f"slide {pos}: trailing space in footer {t!r}")
                if not re.search(r"Slide\. No: \d+$", t):
                    issues.append(f"slide {pos}: footer format odd: {t!r}")

# 2. Irrigation content gone, new papers present -------------------------
full = "\n".join(texts(s) for s in slides)
for bad in ("Irrigation", "irrigation", "Soil Moisture", "ESP8266", "Rice Fields"):
    if bad in full:
        issues.append(f"leftover irrigation content: {bad!r}")
for good in ("Retrieval-Augmented Generation", "Self-RAG", "Hallucination",
             "GraphRAG", "DSPy"):
    if good not in full:
        issues.append(f"expected new paper missing: {good!r}")

# 3. Advanced features removed ------------------------------------------
for bad in ("fine-tune", "fine-tuning", "embedding model on", "cross-paper",
            "labeling query-passage"):
    if bad.lower() in full.lower():
        issues.append(f"advanced-feature claim still present: {bad!r}")

# 4. Streamlit swap + feasibility completion + typos ---------------------
if "React + Vite" in full:
    issues.append("'React + Vite' still present")
if "Streamlit" not in full:
    issues.append("'Streamlit' missing after swap")
if "multi-agent frameworks already studied" not in full:
    issues.append("feasibility sentence not completed")
if "research plan.." in full:
    issues.append("'research plan..' typo still present")
if "features.." in full:
    issues.append("'features..' typo still present")
if "KS KISHNAPRIYA" in full:
    issues.append("name 'KS KISHNAPRIYA' not corrected")

# 5. Table row counts ----------------------------------------------------
for pos, want in ((6, 3), (7, 3), (8, 3), (9, 2)):
    tbl = next(sh.table for sh in slides[pos - 1].shapes
               if getattr(sh, "has_table", False) and sh.has_table)
    got = len(list(tbl.rows))
    if got != want:
        issues.append(f"slide {pos}: table has {got} rows, expected {want}")

# 6. Slide 18 has footer now ---------------------------------------------
if "Slide. No" not in texts(slides[17]):
    issues.append("slide 18: footer missing")

print(f"slides: {len(slides)}")
print("ISSUES:" if issues else "ALL CHECKS PASSED")
for i in issues:
    print(" -", i)

# Print footer strings as they will appear
print("\nFooter strings per slide:")
for pos in range(2, 19):
    for sh in slides[pos - 1].shapes:
        if sh.has_text_frame and "Slide. No" in sh.text_frame.text:
            print(f"  s{pos}: {sh.text_frame.text!r}")
