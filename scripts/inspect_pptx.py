"""Dump the full structure and text of the review PPTX for auditing."""
import sys
from pptx import Presentation
from pptx.util import Emu

PPTX = r"C:\Users\athul\Downloads\FIRST REVIEW new.pptx"


def shape_text(shape, indent="  "):
    lines = []
    kind = shape.shape_type
    name = shape.name
    ph = ""
    if shape.is_placeholder:
        ph = f" [placeholder idx={shape.placeholder_format.idx} type={shape.placeholder_format.type}]"
    lines.append(f"{indent}- {name} ({kind}){ph}")
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            txt = "".join(r.text for r in para.runs)
            if txt.strip():
                lines.append(f"{indent}    L{para.level}: {txt!r}")
    if getattr(shape, "has_table", False) and shape.has_table:
        tbl = shape.table
        for r_i, row in enumerate(tbl.rows):
            cells = []
            for cell in row.cells:
                cells.append(cell.text.replace("\n", " | "))
            lines.append(f"{indent}    row{r_i}: {cells}")
    if shape.shape_type == 6:  # GROUP
        for sub in shape.shapes:
            lines.extend(shape_text(sub, indent + "    "))
    return lines


def main():
    prs = Presentation(PPTX)
    print(f"SLIDE SIZE: {prs.slide_width} x {prs.slide_height} EMU "
          f"({Emu(prs.slide_width).inches:.2f} x {Emu(prs.slide_height).inches:.2f} in)")
    print(f"SLIDE COUNT: {len(prs.slides)}")
    print("=" * 80)

    for i, slide in enumerate(prs.slides, 1):
        print(f"\n########## SLIDE {i} ##########  layout={slide.slide_layout.name!r}")
        for shape in slide.shapes:
            print("\n".join(shape_text(shape)))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                print(f"  NOTES: {notes!r}")

    # Slide masters / layouts footers
    print("\n" + "=" * 80)
    print("MASTERS & LAYOUTS (footer placeholders):")
    for m_i, master in enumerate(prs.slide_masters):
        print(f"\nMASTER {m_i}:")
        for shape in master.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                print(f"  master-shape {shape.name!r}: {shape.text_frame.text!r}")
        for l_i, layout in enumerate(master.slide_layouts):
            texts = [s.text_frame.text for s in layout.shapes
                     if s.has_text_frame and s.text_frame.text.strip()]
            if texts:
                print(f"  layout {l_i} {layout.name!r}: {texts}")


if __name__ == "__main__":
    main()
