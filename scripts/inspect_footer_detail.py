"""Deep inspection: footer box geometry/runs, mojibake chars, table cell paragraphs."""
from pptx import Presentation
from pptx.util import Emu

PPTX = r"C:\Users\athul\Downloads\FIRST REVIEW new.pptx"
prs = Presentation(PPTX)


def emu_in(v):
    return f"{Emu(v).inches:.2f}in" if v is not None else "?"


def dump_footer(slide_idx):
    slide = prs.slides[slide_idx - 1]
    print(f"\n=== SLIDE {slide_idx} footer-ish shapes ===")
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        txt = shape.text_frame.text
        if ("Date" in txt and "IESCE" in txt) or "Slide. No" in txt:
            pos = (f"L={emu_in(shape.left)} T={emu_in(shape.top)} "
                   f"W={emu_in(shape.width)} H={emu_in(shape.height)}")
            print(f"  {shape.name} {pos}")
            print(f"    raw={txt!r}")
            for p_i, para in enumerate(shape.text_frame.paragraphs):
                for r_i, run in enumerate(para.runs):
                    print(f"    p{p_i}r{r_i}: {run.text!r} "
                          f"font={run.font.name} size={run.font.size} "
                          f"bold={run.font.bold}")


for i in (2, 3, 4, 5, 6, 10, 13, 14, 15, 17):
    dump_footer(i)

print("\n=== MOJIBAKE SCAN (U+FFFD) across all slides ===")
for i, slide in enumerate(prs.slides, 1):
    hits = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for m in [c for c in shape.text_frame.text if ord(c) in (0xFFFD,)]:
                hits.append((shape.name, m))
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    bad = [c for c in cell.text if ord(c) == 0xFFFD]
                    if bad:
                        hits.append((f"{shape.name}/cell", len(bad)))
    if hits:
        print(f"  slide {i}: {hits}")

print("\n=== WORK PLAN TABLE (slide 15) exact runs ===")
slide = prs.slides[14]
for shape in slide.shapes:
    if getattr(shape, "has_table", False) and shape.has_table:
        for r_i, row in enumerate(shape.table.rows):
            for cell in row.cells:
                if cell.text.strip():
                    print(f"  row{r_i}: {cell.text!r}")
